"""Verifica la copia local del Excel contra su manifiesto; no descarga ni escribe."""
import json
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.rutas import ROOT, sha256_archivo


def verificar():
    metadata = json.loads((ROOT/'data/raw/procedencia_excel.json').read_text())
    path = ROOT/metadata['archivo_local']
    with path.open('rb') as stream:
        if stream.read(100).startswith(b'version https://git-lfs.github.com/spec/v1'):
            raise ValueError('Hay un puntero Git LFS, no el Excel. Ejecute git lfs pull.')
    if sha256_archivo(path) != metadata['sha256_local']:
        raise ValueError('El Excel no coincide con la fuente congelada; no sobrescribir la versión oficial.')
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name, spec in metadata['hojas'].items():
            if name not in workbook.sheetnames:
                raise ValueError(f'Falta la hoja registrada: {name}')
            sheet = workbook[name]
            if sheet.max_row - 1 != spec['filas'] or sheet.max_column != spec['columnas']:
                raise ValueError(f'Dimensiones distintas de las registradas en la hoja {name}.')
    finally:
        workbook.close()
    return {'resultado': 'correcto', 'sha256_local': metadata['sha256_local'],
            'hojas_verificadas': len(metadata['hojas']), 'licencia': metadata['licencia']['estado']}


if __name__ == '__main__':
    print(json.dumps(verificar(), ensure_ascii=False, indent=2))
